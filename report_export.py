"""Formatted, cautious exports of research source discovery (not AI-generated legal conclusions)."""
from io import BytesIO
from datetime import datetime, timezone
from urllib.parse import urlparse
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.utils import simpleSplit
from xml.sax.saxutils import escape

DISCLAIMER = ('Source discovery only. Records and underlying statements have not been independently verified. '
              'Keyword scores do not establish truth, legal effect, or relevance. Check each original source before relying on it.')

def clean(value, limit=1500):
    return re.sub(r'\s+', ' ', str(value or '')).strip()[:limit]

def collect(data):
    report = data.get('report') or {}
    records = report.get('evidence') or []
    output, seen = [], set()
    for record in records[:300]:
        if not isinstance(record, dict):
            continue
        title = clean(record.get('title') or 'Untitled record', 250)
        url = clean(record.get('url'), 1000)
        key = (url or title.lower())
        if key in seen: continue
        seen.add(key)
        parsed = urlparse(url)
        if parsed.scheme not in ('https','http') or not parsed.netloc: url = ''
        assessment = record.get('evidence_assessment') or {}
        date = clean(record.get('date') or record.get('date_answered'), 40)
        if date: date = date[:10]
        output.append({'title': title, 'source': clean(record.get('source'), 100),
            'date':date, 'url':url, 'description':clean(record.get('description') or record.get('answer'), 600),
            'finding':clean(assessment.get('conclusion'), 350),
            'verified': bool(record.get('verified'))})
    return output

def summary(data, records):
    group = {}
    for r in records:
        s = r['source'] or 'Other source'
        group[s] = group.get(s, 0) + 1
    labels = ', '.join(f'{count} {source}' for source, count in sorted(group.items()))
    return (f'The search identified {len(records)} distinct source '+('lead' if len(records)==1 else 'leads') + (f' ({labels}). ' if labels else '. ') +
        'These are references for investigation, not independently established findings. '
        'The original documents and the legal status of any legislation must be reviewed before substantive conclusions can be drawn.')

def add_doc_link(paragraph, label, url):
    """Insert a normal Word external hyperlink."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    part = paragraph.part
    relation = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), relation)
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '14569A'); props.append(color)
    under = OxmlElement('w:u'); under.set(qn('w:val'), 'single'); props.append(under)
    run.append(props); text = OxmlElement('w:t'); text.text = label; run.append(text); h.append(run)
    paragraph._p.append(h)

def word_report(data):
    question = clean(data.get('question'), 500) or 'Research question'
    items = collect(data)
    d = Document(); sec = d.sections[0]
    sec.header.paragraphs[0].text = 'JOHN SMYTHE  /  RESEARCH ENGINE'
    sec.top_margin = Inches(.8); sec.bottom_margin = Inches(.75)
    normal = d.styles['Normal']; normal.font.name = 'Aptos'; normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    for heading, size in [('Title',23), ('Heading 1',15), ('Heading 2',11)]:
        style = d.styles[heading]; style.font.name = 'Aptos Display'; style.font.size = Pt(size); style.font.color.rgb=RGBColor(21,40,73)
    d.add_heading('RESEARCH BRIEF', 0)
    d.add_paragraph(question, 'Title')
    d.add_paragraph('You ask. We investigate. You decide.')
    d.add_paragraph('Prepared '+ datetime.now(timezone.utc).strftime('%d %B %Y, %H:%M UTC'))
    d.add_heading('Executive summary', 1); d.add_paragraph(summary(data,items))
    d.add_heading('Evidence status and method', 1)
    d.add_paragraph(DISCLAIMER)
    d.add_paragraph('Method: public-source search; records deduplicated by source URL; descriptions and quoted passages are presented as retrieved, without treating them as verified factual conclusions. Dates shown are record dates and may not be the effective dates of laws or policies.')
    d.add_heading('Source findings', 1)
    if not items: d.add_paragraph('No source records were returned. No research conclusion can be established.')
    for idx,r in enumerate(items,1):
        d.add_heading(f'{idx}. {r["title"]}', 2)
        d.add_paragraph('Source: '+(r['source'] or 'Not specified')+'   |   Record date: '+(r['date'] or 'Not provided'))
        if r['description']: d.add_paragraph('Source extract: '+r['description'])
        if r['finding']: d.add_paragraph('Evidence assessment: '+r['finding'])
        if r['url']:
            p=d.add_paragraph('Original record: '); add_doc_link(p,r['url'],r['url'])
        d.add_paragraph('Verification: '+('Marked verified in research data' if r['verified'] else 'Not independently verified'))
    d.add_heading('Research limitations and next steps',1)
    d.add_paragraph('Read the linked primary documents. Confirm statutory commencement, amendments and application where relevant. Compare competing evidence before asserting outcomes or causation. Do not interpret keyword matches as a quality or confidence score.')
    footer = sec.footer.paragraphs[0]; footer.text='JOHN SMYTHE RESEARCH  •  Source discovery — not independently verified'
    b=BytesIO(); d.save(b); return b.getvalue()

def pdf_report(data):
    question=clean(data.get('question'),500) or 'Research question'; items=collect(data)
    b=BytesIO(); pdf=SimpleDocTemplate(b,pagesize=(595.28,841.89), rightMargin=48,leftMargin=48,topMargin=55,bottomMargin=55,
      title='John Smythe Research Brief', author='John Smythe Research Engine')
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='BrandJS',parent=styles['Normal'],textColor=colors.HexColor('#9E2434'),fontSize=10,spaceAfter=15,leading=14))
    styles.add(ParagraphStyle(name='TitleJS',parent=styles['Title'],fontSize=19,leading=24,textColor=colors.HexColor('#152849'),spaceAfter=13))
    styles.add(ParagraphStyle(name='H1JS',parent=styles['Heading2'],fontSize=13,leading=17,spaceBefore=17,spaceAfter=8,textColor=colors.HexColor('#152849')))
    styles.add(ParagraphStyle(name='H2JS',parent=styles['Heading3'],fontSize=10.5,leading=14,spaceBefore=12,spaceAfter=6,textColor=colors.HexColor('#152849')))
    styles.add(ParagraphStyle(name='BodyJS',parent=styles['BodyText'],fontSize=9,leading=14,spaceAfter=8,wordWrap='CJK'))
    story=[]
    def p(t,style='BodyJS'): story.append(Paragraph(escape(clean(t,2600)),styles[style]))
    p('JOHN SMYTHE  /  RESEARCH ENGINE','BrandJS'); p('RESEARCH BRIEF — '+question,'TitleJS')
    p('You ask. We investigate. You decide. | '+datetime.now(timezone.utc).strftime('%d %B %Y, %H:%M UTC'))
    p('Executive summary','H1JS'); p(summary(data,items))
    p('Evidence status and method','H1JS'); p(DISCLAIMER)
    p('Method: public-source search; results deduplicated by source URL. Record dates are not necessarily the effective dates of laws or policies. No independent source-content or legal-status verification is claimed.')
    p('Source findings','H1JS')
    if not items: p('No source records were returned. No research conclusion can be established.')
    for idx,r in enumerate(items,1):
        block=[Paragraph(escape(clean(f'{idx}. {r["title"]}',300)),styles['H2JS']),
               Paragraph(escape(f'Source: {r["source"] or "Not specified"} | Record date: {r["date"] or "Not provided"}'),styles['BodyJS'])]
        story.append(KeepTogether(block))
        if r['description']: p('Source extract: '+r['description'])
        if r['finding']: p('Evidence assessment: '+r['finding'])
        if r['url']:
            u=escape(r['url'], {'"':'&quot;'})
            story.append(Paragraph('Original record: <link href="'+u+'" color="#14569A">'+escape(r['url'])+'</link>',styles['BodyJS']))
        p('Verification: '+('Marked verified in research data' if r['verified'] else 'Not independently verified'))
    p('Research limitations and next steps','H1JS')
    p('Read linked primary documents. Confirm statutory commencement, amendments and application where relevant. Compare evidence before asserting outcomes or causation. Keyword matches are not confidence scores.')
    def page_number(canvas, doc):
        canvas.saveState(); canvas.setFont('Helvetica',8); canvas.setFillColor(colors.HexColor('#596579'))
        canvas.drawString(48,35,'JOHN SMYTHE RESEARCH  •  Unverified source discovery')
        canvas.drawRightString(547,35,str(doc.page)); canvas.restoreState()
    pdf.build(story,onFirstPage=page_number,onLaterPages=page_number)
    return b.getvalue()
